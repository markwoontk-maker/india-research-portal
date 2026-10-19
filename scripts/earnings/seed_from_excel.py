"""One-time seed of data/earnings.json from the analyst's earnings workbook.

Run:
  uv run --python 3.12 --with openpyxl python -I scripts/earnings/seed_from_excel.py "<path to .xlsx>"

Reads Sheet2 of the workbook (layout: A=sector, B=company, C-E revenue YoY
1Q27-3Q27, F-H EBITDA YoY 1Q27-3Q27, I=metric label / outlook note,
J-P revenue 1Q26-3Q27, Q-W EBITDA 1Q26-3Q27) and writes every company row
into data/earnings.json. Cell fills carry the beat/miss call (theme accent6
green / FF92D050 = beat, theme accent2 orange = miss) and cell comments carry
the explanation that the page shows as a pop-up. Rows the miner already
filled from a broker PDF (source.kind == "pdf") are never overwritten.
"""
import json, os, re, sys
from datetime import date

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, 'data', 'earnings.json')
COMPANIES = json.load(open(os.path.join(ROOT, 'data', 'companies.json'), encoding='utf-8'))
LIB = r'C:\Users\admin\Desktop\India Related Reports'

QTRS_ABS = ['1Q26', '2Q26', '3Q26', '4Q26', '1Q27', '2Q27', '3Q27']
QTRS_YOY = ['1Q27', '2Q27', '3Q27']

# Workbook name -> canonical company name (companies.json key / library folder).
ALIAS = {
    'Eicher': 'Eicher Motors', 'TVS Motors': 'TVS Motor', 'Hyundai India': 'Hyundai Motor India',
    'HFCL Ltd': 'HFCL', 'Sterlite Tech': 'Sterlite Tech', 'Interglobe Aviation': 'InterGlobe Aviation',
    'Allied Blenders': 'Allied Blenders & Distillers', 'United Spirits': 'United Spirits',
    'Billionbrains Garage Ventures': "Billionbrains Garage Ventures ('Groww')", 'ICICI AMC': 'ICICI Prudential AMC',
    'Kfin Tech': 'KFin Technologies', 'Nippon AMC': 'Nippon AMC', 'Nuvama': 'Nuvama',
    'Exide Indutries': 'Exide Industries', 'Sona BLW': 'Sona BLW', 'Canfin': 'Can Fin Homes',
    'HDFC Asset Management': 'HDFC AMC', 'Kotak Mahindra Bank': 'Kotak Mahindra Bank',
    'One 97 Communications': 'One 97 Communications', 'Blue Jet Healthcare': 'Blue Jet Healthcare',
    "Divi's": "Divi's Laboratories", 'Ambuja': 'Ambuja Cements', 'Ultratech Cement': 'UltraTech Cement',
    'Kansai Nerolac Paints': 'Kansai Nerolac Paints', 'UPL Limited': 'UPL', 'Adani Enterprise': 'Adani Enterprises',
    'Blackbuck Ltd': 'Blackbuck', 'IRB Infra': 'IRB Infrastructure', 'Dixon': 'Dixon Technologies',
    'Syrma SGS Technology': 'Syrma SGS Technology', "Westlife (McDonald's)": 'Westlife Foodworld',
    'Central Depository Services': 'Central Depository Services (India)', 'Dr Lal Pathlabs': 'Dr Lal PathLabs',
    'Indian Kotels': 'Indian Hotels Company', 'Leela Palaces (Luxury)': 'Leela Palaces Hotels & Resorts',
    'ICICI Life': 'ICICI Prudential Life', 'Niva Bupa': 'Niva Bupa Health Insurance', 'Star Health': 'Star Health',
    'Bharti Hexaware': 'Hexaware Technologies', 'Fractal Analysis': 'Fractal Analytics', 'HCL Tech': 'HCL Technologies',
    'LTM': 'LTIMindtree', 'Tata Consultancy Services': 'Tata Consultancy Services', 'Container Corp of India': 'Container Corporation of India',
    'Aptus': 'Aptus Value Housing Finance', 'Cholamandalam': 'Cholamandalam', 'Home First': 'Home First Finance',
    'LIC Housing': 'LIC Housing Finance', 'Mahindra Finance': 'Mahindra Finance', 'Muthoot': 'Muthoot Finance',
    'Gail India (LNG, LPG)': 'GAIL India', 'Indian Oil Corp': 'Indian Oil Corporation', "Dr Reddy's": "Dr Reddy's Laboratories",
    'Mankind': 'Mankind Pharma', 'Sun Pharma': 'Sun Pharmaceuticals', 'Torrent Pharma': 'Torrent Pharmaceuticals',
    'Adani Ports': 'Adani Ports and SEZ', 'Bharat Heavy Electricals': 'Bharat Heavy Electricals', 'Thermax Limited': 'Thermax',
    'Warree Energies': 'Waaree Energies', 'Embassy REIT': 'Embassy Office Parks REIT', 'Lodha': 'Lodha Developers',
    'Prestige Estates': 'Prestige Estates Projects', 'Supreme Industries': 'Supreme Industries', 'Emmvee': 'Emmvee Photovoltaic Power',
    'Avenue Supermart': 'Avenue Supermarts', 'Colgate': 'Colgate-Palmolive (India)', 'Dabur': 'Dabur India',
    'Nestle India': 'Nestle India', 'Tata Consumer': 'Tata Consumer Products', 'Jindal Steel (Long Steel)': 'Jindal Steel',
    'Shyam Metalics & Energy': 'Shyam Metalics and Energy', 'Indus Tower': 'Indus Towers', 'Kajaria Ceramics Ltd': 'Kajaria Ceramics',
    'Adani Green': 'Adani Green Energy', 'Bajaj Finserv': 'Bajaj Finserv', 'Go Digit': 'Go Digit General Insurance',
    'ICICI Lombard General Insurance': 'ICICI Lombard General Insurance', 'Tbo Tek': 'TBO Tek', 'CRISIL': 'CRISIL',
    'Global Health (Medanta)': 'Global Health (Medanta)', 'Narayana Hrudayalaya': 'Narayana Hrudayalaya',
    'Vedanta Aluminium': 'Vedanta Aluminium', 'Piramal Finance': 'Piramal Finance', 'Steel Authority of India': 'Steel Authority of India',
}


METRIC_RE = re.compile(r'\b(PPOP|EBIT|EBITDA|PBT|PAT|APE|VNB|NEP|GWP|GDPI|NII|Consolidated|Standalone)\b|US\$', re.I)


def metric_labels(sector, basis):
    """Revenue / EBITDA labels for a row: lenders use Total operating income /
    PPOP, insurers GWP / PBT (life: APE / VNB per the sheet), IT shows EBIT
    when the sheet says so."""
    s, b = sector.strip().lower(), basis.upper()
    rev, ebd = 'Revenue', 'EBITDA'
    if s in ('banks', 'nbfc'):
        rev, ebd = 'Total op. income', 'PPOP'
    elif s == 'insurance':
        rev, ebd = 'GWP', 'PBT'
    if '/' in b:
        a, c = [x.strip() for x in b.split('/', 1)]
        if a and c:
            return {'rev': a, 'ebitda': c}
    for tok in ('PPOP', 'EBITDA', 'EBIT', 'PBT', 'PAT'):
        if re.search(r'\b' + tok + r'\b', b):
            ebd = tok
            break
    return {'rev': rev, 'ebitda': ebd}


def norm(s):
    s = re.sub(r"\(.*?\)", ' ', str(s).lower())
    s = re.sub(r'\b(ltd|limited|india|the|co|company|industries|inds)\b', ' ', s)
    return re.sub(r'[^a-z0-9]', '', s)


def known_names():
    names = set(COMPANIES)
    try:
        names |= {d for d in os.listdir(LIB) if os.path.isdir(os.path.join(LIB, d))}
    except OSError:
        pass
    return names


def resolve(raw, names):
    raw = raw.strip()
    if raw in ALIAS:
        return ALIAS[raw]
    if raw in names:
        return raw
    k = norm(raw)
    hits = [n for n in names if norm(n) == k]
    if hits:
        return sorted(hits, key=lambda n: (n not in COMPANIES, len(n)))[0]
    hits = [n for n in names if k and (norm(n).startswith(k) or k.startswith(norm(n)) and len(norm(n)) > 4)]
    if len(hits) == 1:
        return hits[0]
    return raw


def fill_call(cell):
    f = cell.fill
    if not f or f.fill_type != 'solid':
        return None
    c = f.fgColor
    if c.type == 'theme':
        return {9: 'beat', 5: 'miss'}.get(c.theme)
    if c.type == 'rgb' and isinstance(c.rgb, str):
        return {'FF92D050': 'beat', 'FF00B050': 'beat', 'FFFF0000': 'miss', 'FFFFC7CE': 'miss', 'FFC6EFCE': 'beat'}.get(c.rgb.upper())
    return None


def comment(cell):
    if not cell.comment:
        return None
    t = cell.comment.text
    t = re.sub(r'^[^:\n]{0,40}:\n', '', t)          # drop "Author:\n" header if present
    t = t.replace('\ufffd', "'").replace('\r', '')
    t = re.sub(r'\s*\n\s*', ' ', t).strip()
    return t or None


def num(v):
    if isinstance(v, bool) or v is None or v == '':
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def main(path):
    wb_f = openpyxl.load_workbook(path)                  # formulas + styles + comments
    wb_v = openpyxl.load_workbook(path, data_only=True)  # cached values
    ws, wv = wb_f['Sheet2'], wb_v['Sheet2']
    names = known_names()

    existing, skipped = {}, {}
    if os.path.exists(OUT):
        old_file = json.load(open(OUT, encoding='utf-8'))
        existing = {c['name']: c for c in old_file.get('companies', [])}
        skipped = old_file.get('skipped', {})  # miner's "not a results note" memory

    rows, seen = [], set()
    for r in range(3, ws.max_row + 1):
        raw = ws.cell(r, 2).value
        if not raw or not str(raw).strip():
            continue
        name = resolve(str(raw), names)
        if name in seen:          # duplicate row (e.g. Swiggy twice) -> keep the richer one below
            prev = next(x for x in rows if x['name'] == name)
            rows.remove(prev)
        seen.add(name)
        sector = str(ws.cell(r, 1).value or '').strip()
        label = ws.cell(r, 9).value
        label = str(label).strip() if label else ''
        is_basis = bool(label) and len(label) <= 32 and bool(METRIC_RE.search(label))
        basis = label if is_basis else ''
        note = '' if is_basis else label.replace('\ufffd', "'")

        q = {}
        for i, qt in enumerate(QTRS_ABS):
            rev, ebd = num(wv.cell(r, 10 + i).value), num(wv.cell(r, 17 + i).value)
            if rev is not None or ebd is not None:
                q.setdefault(qt, {})
                if rev is not None: q[qt]['rev'] = rev
                if ebd is not None: q[qt]['ebitda'] = ebd
            for col, key in ((10 + i, 'revNote'), (17 + i, 'ebitdaNote')):
                t = comment(ws.cell(r, col))
                if t: q.setdefault(qt, {})[key] = t
        def yoy(cell, k, qt):
            # Hard-typed growth numbers are taken as-is. Formula cells (=N/J-1) are
            # recomputed from the absolutes: Excel treats a blank current quarter as 0,
            # so its cached value is a bogus -100% for quarters not reported yet.
            if not (isinstance(cell.value, str) and cell.value.startswith('=')):
                return num(cell.value)
            prev = qt[:2] + str(int(qt[2:]) - 1)
            a, b = q.get(qt, {}).get(k), q.get(prev, {}).get(k)
            return round(a / b - 1, 4) if a is not None and b and b > 0 else None

        for i, qt in enumerate(QTRS_YOY):
            rc, ec = ws.cell(r, 3 + i), ws.cell(r, 6 + i)
            ry, ey = yoy(rc, 'rev', qt), yoy(ec, 'ebitda', qt)
            for cell, val, k in ((rc, ry, 'rev'), (ec, ey, 'ebitda')):
                call, note_c = fill_call(cell), comment(cell)
                if val is None and not call and not note_c:
                    continue
                d = q.setdefault(qt, {})
                if val is not None: d[k + 'YoY'] = round(val, 4)
                if call: d[k + 'Vs'] = call
                elif val is not None: d[k + 'Vs'] = 'inline'
                if note_c: d[k + 'Note'] = (d.get(k + 'Note', '') + ' ' + note_c).strip() if d.get(k + 'Note') else note_c
        # compute any YoY the sheet left as a blank formula but whose absolutes exist
        for qt in QTRS_YOY:
            prev = 'Q'.join([qt[:1], str(int(qt[2:]) - 1)])
            for k in ('rev', 'ebitda'):
                a, b = q.get(qt, {}).get(k), q.get(prev, {}).get(k)
                if a is not None and b and b > 0 and qt in q and q[qt].get(k + 'YoY') is None:
                    q[qt][k + 'YoY'] = round(a / b - 1, 4)
                    q[qt].setdefault(k + 'Vs', 'inline')

        reported = [qt for qt in QTRS_YOY if q.get(qt, {}).get('revYoY') is not None or q.get(qt, {}).get('ebitdaYoY') is not None]
        overall = {}
        if reported:
            latest = reported[-1]
            call = fill_call(ws.cell(r, 2))
            if not call:
                calls = [q[latest].get('revVs'), q[latest].get('ebitdaVs')]
                b, m = calls.count('beat'), calls.count('miss')
                call = 'beat' if b > m else 'miss' if m > b else 'inline'
            for qt in reported:
                if qt == latest:
                    q[qt]['overall'] = call
                elif 'overall' not in q[qt]:
                    calls = [q[qt].get('revVs'), q[qt].get('ebitdaVs')]
                    b, m = calls.count('beat'), calls.count('miss')
                    q[qt]['overall'] = 'beat' if b > m else 'miss' if m > b else 'inline'
            for qt in reported:
                q[qt]['src'] = {'kind': 'sheet', 'file': os.path.basename(path)}

        row = {'name': name, 'sheetName': str(raw).strip(), 'sector': sector,
               'ticker': (COMPANIES.get(name) or {}).get('ticker', ''),
               'basis': basis, 'note': note, 'unit': None, 'revUnit': None,
               'metrics': metric_labels(sector, basis), 'q': q}
        if sector == 'IT':  # IT revenue is tracked in US$; the sheet's small numbers are US$ mn
            last = [d['rev'] for _, d in sorted(q.items()) if d.get('rev')]
            if last and last[-1] < 15000:
                row['revUnit'] = 'US$ mn'
        old = existing.get(name)
        if old:  # never clobber quarters the PDF miner already wrote
            for qt, d in (old.get('q') or {}).items():
                if (d.get('src') or {}).get('kind') == 'pdf':
                    row['q'][qt] = d
            for k in ('unit', 'revUnit', 'sym', 'mcap', 'mcapAsOf'):
                if old.get(k): row[k] = old[k]
        rows.append(row)

    infer_units(rows)
    for name, c in existing.items():  # keep companies the miner added that aren't in the sheet
        if name not in seen:
            rows.append(c)
    rows.sort(key=lambda c: (c['sector'].lower(), c['name'].lower()))
    out = {'updated': date.today().isoformat(), 'source': 'Analyst workbook (Sheet2) + broker results notes in the PDF library',
           'quarters': QTRS_ABS, 'companies': rows, 'skipped': skipped}
    with open(OUT, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)
    unmatched = [c['sheetName'] for c in rows if c.get('sheetName') and c['name'] not in names]
    print(f'wrote {len(rows)} companies -> {OUT}')
    print('names not found in companies.json / library:', unmatched)


def infer_units(rows):
    """Guess each row's unit by comparing 4 x the latest quarterly revenue with
    FY26 revenue in data/financials.json (which is in INR mn)."""
    try:
        fin = json.load(open(os.path.join(ROOT, 'data', 'financials.json'), encoding='utf-8'))
    except (OSError, ValueError):
        return
    scales = [(1, 'INR mn'), (0.1, 'INR cr'), (0.001, 'INR bn'), (10, 'INR lakh')]
    for c in rows:
        if c.get('unit'):
            continue
        f = fin.get(c['name']) or {}
        fy = (((f.get('metrics') or {}).get('revenue') or {}).get('FY26') or {}).get('avg')
        vals = [d.get('rev') for qt, d in sorted(c['q'].items()) if d.get('rev')]
        if not fy or not vals:
            continue
        ratio = vals[-1] * 4 / fy
        best = min(scales, key=lambda s: abs(__import__('math').log(ratio / s[0])))
        if abs(__import__('math').log(ratio / best[0])) < 0.9:
            c['unit'] = best[1]


if __name__ == '__main__':
    main(sys.argv[1])
